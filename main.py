from docx import Document
from docx.shared import RGBColor
import re
import time
import os
from session import is_logged_in

def fill_invitation(template_path, output_path, data):
    # Cek apakah user sudah login (opsional - jika mau protect fungsi ini)
    # if not is_logged_in():
    #     raise Exception("Unauthorized: Please login first!")
    
    doc = Document(template_path)

    def replace_in_paragraph(paragraph):
        """Ganti placeholder sambil mempertahankan formatting"""
        full_text = paragraph.text
        
        needs_replacement = False
        for key in data.keys():
            pattern = re.compile(re.escape(key), re.IGNORECASE)
            if pattern.search(full_text):
                needs_replacement = True
                break
        
        if not needs_replacement:
            return
        
        if not paragraph.runs:
            return
        
        default_format = paragraph.runs[0]
        
        new_text = full_text
        for key, value in data.items():
            pattern = re.compile(re.escape(key), re.IGNORECASE)
            new_text = pattern.sub(value, new_text)
        
        for _ in range(len(paragraph.runs)):
            paragraph._element.remove(paragraph.runs[0]._element)
        
        new_run = paragraph.add_run(new_text)
        new_run.bold = default_format.bold
        new_run.italic = default_format.italic
        new_run.underline = default_format.underline
        new_run.font.size = default_format.font.size
        new_run.font.name = default_format.font.name
        if default_format.font.color.rgb:
            new_run.font.color.rgb = default_format.font.color.rgb

    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            replace_in_paragraph(paragraph)
        for paragraph in section.footer.paragraphs:
            replace_in_paragraph(paragraph)

    timestamp_human = time.strftime("%Y-%m-%d %H:%M:%S")
    doc.add_paragraph(f"Dibuat otomatis pada: {timestamp_human}")

    ts = time.strftime("%Y%m%d_%H%M%S")
    dir_name, file_name = os.path.split(output_path)
    stem, ext = os.path.splitext(file_name)
    final_name = f"{stem}_{ts}{ext}"
    final_path = os.path.join(dir_name, final_name) if dir_name else final_name

    doc.save(final_path)
    return final_path

if __name__ == "__main__":
    # Proteksi: main.py tidak bisa dijalankan langsung
    if not is_logged_in():
        print("❌ Error: Anda harus login terlebih dahulu!")
        print("Jalankan: python3 login.py")
        exit(1)
    
    # Contoh test jika sudah login
    data = {
        '[date]': '5 Januari 2025',
        '[perihal]': 'Undangan Rapat Koordinasi',
        '[Nama_Panjang]': 'Siti Aminah',
        '[nama_panjang]': 'Siti Aminah'
    }
    template_path = 'template.docx'
    output_path = 'output.docx'
    fill_invitation(template_path, output_path, data)